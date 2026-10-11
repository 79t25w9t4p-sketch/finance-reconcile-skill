#!/usr/bin/env python3
"""Reproduce a fictional single-customer month-end reconciliation."""
import argparse
import json
from pathlib import Path
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from doctor import check_environment

LEFT = [
    ['SH-001', '1200.00', '模拟：普通送货'],
    ['SH-002', '2350.00', '模拟：两表金额不同'],
    ['SH-003', '800.00', '模拟：普通送货'],
    ['SH-004', '600.00', '模拟：只在我方表出现'],
    ['SH-005', '450.00', '模拟：同单号第一条'],
    ['SH-005', '450.00', '模拟：同单号第二条，需核实粒度'],
    ['SH-006', '', '模拟：金额缺失，不能按零处理'],
    ['SH-007', '1500.00', '模拟：普通送货'],
]
RIGHT = [
    ['SH-007', '1500.00', '模拟：顺序不同'],
    ['SH-001', '1200.00', '模拟：普通送货'],
    ['SH-002', '2300.00', '模拟：两表金额不同'],
    ['SH-003', '800.00', '模拟：普通送货'],
    ['SH-005', '450.00', '模拟：需核实是否一单多行'],
    ['SH-006', '900.00', '模拟：我方金额缺失'],
    ['SH-008', '600.00', '模拟：只在客户表出现'],
]


def write_input(path, headers, rows):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '明细'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(bold=True, color='FFFFFF')
    for col, width in [('A', 24), ('B', 20), ('C', 48)]:
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description='用虚构月末客户明细生成首次复核与跟进清单')
    parser.add_argument('out', help='新的演示目录，不覆盖已有文件')
    parser.add_argument('--node', default='node')
    args = parser.parse_args()
    report = check_environment(args.node)
    if not report['ready']:
        for item in report['checks']:
            if not item['ok']:
                print(item['name']+'：'+item['detail']+'；'+item['next'], file=sys.stderr)
        return 2
    dest = Path(args.out).resolve()
    if dest.exists():
        raise ValueError('演示目录已存在，请使用新目录')
    with tempfile.TemporaryDirectory(dir=dest.parent) as scratch:
        stage = Path(scratch)/'demo'
        stage.mkdir()
        left = stage/'模拟_我方销售明细.xlsx'
        right = stage/'模拟_客户采购明细.xlsx'
        write_input(left, ['送货单号', '含税金额', '说明'], LEFT)
        write_input(right, ['供方送货单号', '含税金额', '说明'], RIGHT)
        command = [sys.executable, str(ROOT/'scripts/workflow.py'), '--node', args.node,
                   '--left', str(left), '--right', str(right),
                   '--left-keys', '["送货单号"]', '--right-keys', '["供方送货单号"]',
                   '--left-amount', '含税金额', '--right-amount', '含税金额',
                   '--save-config', str(stage/'月末客户核对配置.json'), '--out', str(stage/'核对结果')]
        result = subprocess.run(command, capture_output=True, text=True, timeout=60)
        if result.returncode:
            raise ValueError(result.stderr.strip())
        summary = json.loads((stage/'核对结果/摘要.json').read_text(encoding='utf-8'))
        (stage/'示例口径.txt').write_text(
            '全部为虚构数据，仅作操作演示。假定同一客户、2026年9月、人民币、'
            '每张送货单一行、同一含税金额口径。SH-005故意重复，需人工核实。'
            'SH-006故意缺失金额，不能按零处理。不是实际客户案例或耗时测量。\n', encoding='utf-8')
        stage.rename(dest)
    print(json.dumps({'demo': str(dest), 'leftRows': summary['leftRows'],
                      'rightRows': summary['rightRows'], 'counts': summary['counts'],
                      'workbook': str(dest/'核对结果/复核工作簿.xlsx'),
                      'followup': str(dest/'核对结果/跟进清单.xlsx')}, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as error:
        print('示例未完成：'+str(error), file=sys.stderr)
        sys.exit(1)
