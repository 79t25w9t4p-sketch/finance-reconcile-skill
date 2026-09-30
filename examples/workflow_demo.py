"""Create two small fictional runs to demonstrate review carry-forward."""
from pathlib import Path
import os
import subprocess
import sys
import openpyxl

root = Path(__file__).resolve().parents[1]
dest = Path(sys.argv[1])
dest.mkdir(parents=True, exist_ok=False)
left = dest/'左表.csv'
right = dest/'右表.csv'
left.write_text('编号,金额\nA,100\nB,200\nC,300', encoding='utf8')
right.write_text('编号,金额\nA,99\nB,199\nC,299', encoding='utf8')
base = [sys.executable, str(root/'scripts/workflow.py'), '--node', os.environ.get('NODE','node'), '--left', str(left), '--right', str(right)]
subprocess.run(base+['--save-config',str(dest/'核对配置.json'),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--out',str(dest/'首次')], check=True)
prior=dest/'首次/复核工作簿.xlsx'
wb=openpyxl.load_workbook(prior)
for row in range(2,5):
    wb['待处理'].cell(row,4,'已解决')
    wb['待处理'].cell(row,5,'模拟意见：已核实此项，仅供演示')
wb.save(prior)
left.write_text('编号,金额\nB,201\nA,100\nD,400', encoding='utf8')
right.write_text('编号,金额\nA,99\nB,199\nD,399', encoding='utf8')
subprocess.run(base+['--config',str(dest/'核对配置.json'),'--previous',str(prior),'--out',str(dest/'重跑')], check=True)
