import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import prepare as module
from reconcile import fingerprint, load_table


class Preparation(unittest.TestCase):
    def workbook(self, path, rows, sheet='明细'):
        workbook = openpyxl.Workbook()
        workbook.active.title = sheet
        for row in rows:
            workbook.active.append(row)
        workbook.save(path)
        workbook.close()

    def apply(self, path, out, **kwargs):
        options = dict(expected_sha256=fingerprint(path), exclude_rows=[3],
                       confirmed=True, reason='模拟确认：第3行是汇总，非逐单明细', out=out)
        options.update(kwargs)
        return module.prepare(path, **options)

    def test_inspection_lists_all_positions_without_values_or_writes(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'export.csv'
            path.write_text('编号,金额\nA,123.45\n' + '合计,123.45\n' * 25, encoding='utf-8')
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/prepare.py'), '--input', str(path), '--inspect'],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            info = json.loads(result.stdout)
            self.assertEqual(info['source']['sha256'], fingerprint(path))
            self.assertEqual(info['dataWarnings'][0]['rows'], list(range(3, 28)))
            self.assertTrue(info['dataWarnings'][0]['positionsComplete'])
            self.assertNotIn('123.45', result.stdout)
            self.assertEqual(list(Path(folder).iterdir()), [path])

    def test_xlsx_selected_page_header_and_original_positions_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'export.xlsx'
            self.workbook(path, [['标题'], [], ['编号', '金额'], ['0001', '100'], [],
                                 ['编号', '金额'], [], ['0002', '50'], ['合计', '150']])
            wb = openpyxl.load_workbook(path)
            wb.create_sheet('说明')['A1'] = '=1+1'
            wb.save(path); wb.close()
            before = fingerprint(path)
            result = self.apply(path, Path(folder) / 'copy', sheet='明细', header_row='3', exclude_rows=[6, 9])
            self.assertEqual(result['copy']['rows'], 2)
            table, _ = load_table(Path(folder) / 'copy/核对副本.xlsx', '明细', '3')
            self.assertEqual(table['sourceLines'], [4, 8])
            self.assertEqual(table['rows'], [['0001', '100'], ['0002', '50']])
            wb = openpyxl.load_workbook(Path(folder) / 'copy/核对副本.xlsx')
            self.assertEqual(wb.sheetnames, ['明细'])
            self.assertIsNone(wb['明细']['A6'].value)
            wb.close()
            audit = json.loads((Path(folder) / 'copy/整理说明.json').read_text(encoding='utf-8'))
            self.assertEqual(audit['source']['sha256'], before)
            self.assertEqual(audit['retainedPositions'], [4, 8])
            self.assertEqual(audit['copy']['sha256'], fingerprint(Path(folder) / 'copy/核对副本.xlsx'))
            self.assertEqual(fingerprint(path), before)

    def test_confirmation_and_current_candidate_guards_leave_no_output(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\nA,10\n合计,10\n', encoding='utf-8')
            before = fingerprint(path)
            invalid = [{'confirmed': False}, {'reason': ' '}, {'expected_sha256': None},
                       {'exclude_rows': []}, {'exclude_rows': [True]}, {'exclude_rows': [3, 3]},
                       {'exclude_rows': [1]}, {'exclude_rows': [2]}, {'exclude_rows': [99]},
                       {'exclude_rows': [3], 'exclude_kinds': ['summary_rows']},
                       {'exclude_rows': None, 'exclude_kinds': ['repeated_headers']}]
            for index, options in enumerate(invalid):
                out = Path(folder) / str(index)
                with self.subTest(options=options), self.assertRaises(ValueError):
                    self.apply(path, out, **options)
                self.assertFalse(out.exists())
            self.assertEqual(fingerprint(path), before)

    def test_stale_hash_cannot_reuse_old_row_positions(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\nA,10\n合计,10\n', encoding='utf-8')
            old = fingerprint(path)
            path.write_text('编号,金额\nA,20\n合计,20\n', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, '输入已变化'):
                self.apply(path, Path(folder) / 'copy', expected_sha256=old)
            self.assertFalse((Path(folder) / 'copy').exists())

    def test_confirmed_whole_group_handles_more_than_twenty_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\nTOTAL-001,10\n合计-001,20\n' + '合计,30\n' * 25, encoding='utf-8')
            result = self.apply(path, Path(folder) / 'copy', exclude_rows=None, exclude_kinds=['summary_rows'])
            self.assertEqual(result['excludedCount'], 25)
            self.assertEqual(result['remainingWarnings'], [])
            table, _ = load_table(Path(folder) / 'copy/核对副本.xlsx')
            self.assertEqual(table['rows'], [['TOTAL-001', '10'], ['合计-001', '20']])
            self.assertEqual(table['sourceLines'], [2, 3])

    def test_partial_confirmation_does_not_remove_unconfirmed_label(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\n合计,10\n编号,金额\n', encoding='utf-8')
            result = self.apply(path, Path(folder) / 'copy', reason='模拟确认：只排除第3行重复表头')
            self.assertEqual(result['remainingWarnings'][0]['rows'], [2])
            table, _ = load_table(Path(folder) / 'copy/核对副本.xlsx')
            self.assertEqual(table['rows'], [['合计', '10']])

    def test_csv_tsv_conversion_preserves_text_and_logical_record_provenance(self):
        with tempfile.TemporaryDirectory() as folder:
            for suffix, delimiter in [('csv', ','), ('tsv', '\t')]:
                path = Path(folder) / ('source.' + suffix)
                rows = ['编号' + delimiter + '金额' + delimiter + '备注',
                        '0001' + delimiter + '10' + delimiter + '"第一行\n第二行"',
                        '合计' + delimiter + '30' + delimiter,
                        '=TEXT-ID' + delimiter + '20' + delimiter + '"=HYPERLINK(""fake"")"']
                # A quoted formula-like value is input text; it must stay text in XLSX.
                path.write_bytes(('\n'.join(rows) + '\n').encode('utf-8'))
                out = Path(folder) / suffix
                self.apply(path, out)
                table, _ = load_table(out / '核对副本.xlsx')
                self.assertEqual(table['sourceLines'], [2, 4])
                self.assertEqual(table['rows'][0], ['0001', '10', '第一行\n第二行'])
                self.assertEqual(table['rows'][1], ['=TEXT-ID', '20', '=HYPERLINK("fake")'])
                audit = json.loads((out / '整理说明.json').read_text(encoding='utf-8'))
                self.assertEqual(audit['positionKind'], 'parsed_record')
                wb = openpyxl.load_workbook(out / '核对副本.xlsx')
                self.assertFalse(any(c.data_type == 'f' for row in wb.active for c in row))
                wb.close()

    def test_crlf_and_cr_text_survive_without_optional_lxml(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); path = root / 'source.csv'; out = root / 'copy'
            path.write_bytes('编号,金额,备注\r\nA,10,"一\r\n二\r三\n四"\r\n合计,10,\r\n'.encode('utf-8'))
            before = fingerprint(path)
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/prepare.py'), '--input', str(path),
                                     '--expect-sha256', before, '--exclude-rows', '[3]', '--confirmed',
                                     '--reason', '模拟确认：第3行汇总', '--out', str(out)],
                                    capture_output=True, text=True, env={**os.environ, 'OPENPYXL_LXML': 'False'})
            self.assertEqual(result.returncode, 0, result.stderr)
            table, _ = load_table(out / '核对副本.xlsx')
            self.assertEqual(table['rows'], [['A', '10', '一\r\n二\r三\n四']])
            self.assertEqual(table['sourceLines'], [2])
            self.assertEqual(fingerprint(path), before)

    def test_formula_source_cannot_be_cleaned_away(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            self.workbook(path, [['编号', '金额'], ['A', '10'], ['合计', '=SUM(B2)']])
            with self.assertRaisesRegex(ValueError, 'Formula'):
                self.apply(path, Path(folder) / 'copy')
            self.assertFalse((Path(folder) / 'copy').exists())

    def test_source_change_during_preparation_leaves_no_copy(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\nA,10\n合计,10\n', encoding='utf-8')
            original = module.load_table
            def change_after_copy(file, *args, **kwargs):
                result = original(file, *args, **kwargs)
                if Path(file).name == '核对副本.xlsx':
                    path.write_text('编号,金额\nA,20\n合计,20\n', encoding='utf-8')
                return result
            with patch.object(module, 'load_table', side_effect=change_after_copy):
                with self.assertRaisesRegex(ValueError, '整理过程中输入发生变化'):
                    self.apply(path, Path(folder) / 'copy')
            self.assertFalse((Path(folder) / 'copy').exists())
            self.assertEqual(sorted(p.name for p in Path(folder).iterdir()), ['source.csv'])

    def test_existing_output_and_empty_detail_are_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.csv'
            path.write_text('编号,金额\nA,10\n合计,10\n', encoding='utf-8')
            out = Path(folder) / 'copy'; out.mkdir()
            marker = out / '原有文件.txt'; marker.write_text('keep', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, '输出目录已存在'):
                self.apply(path, out)
            self.assertEqual(marker.read_text(encoding='utf-8'), 'keep')
            path.write_text('编号,金额\n合计,10\n', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, '排除后没有明细'):
                self.apply(path, Path(folder) / 'empty', exclude_rows=[2])
            self.assertFalse((Path(folder) / 'empty').exists())

    def test_prepared_copy_runs_workflow_with_original_csv_config(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            left = root / 'left.csv'; right = root / 'right.csv'
            left.write_text('编号,金额\nA,100\nB,50\n', encoding='utf-8')
            right.write_text('编号,金额\nA,99\nB,50\n', encoding='utf-8')
            command = [sys.executable, str(ROOT / 'scripts/workflow.py'), '--left', str(left), '--right', str(right)]
            first = subprocess.run(command + ['--left-keys', '["编号"]', '--right-keys', '["编号"]',
                                   '--left-amount', '金额', '--right-amount', '金额', '--save-config', str(root / 'config.json'),
                                   '--out', str(root / 'first')], capture_output=True, text=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            config = (root / 'config.json').read_bytes()
            previous = root / 'first/复核工作簿.xlsx'
            wb = openpyxl.load_workbook(previous)
            wb['待处理']['D2'] = '已解决'
            wb['待处理']['E2'] = '模拟确认：差额1已核实'
            wb.save(previous); wb.close()
            left.write_text('编号,金额\nA,100\n编号,金额\nB,50\n合计,150\n', encoding='utf-8')
            self.apply(left, root / 'copy', exclude_rows=[3, 5])
            second = subprocess.run([sys.executable, str(ROOT / 'scripts/workflow.py'), '--left', str(root / 'copy/核对副本.xlsx'),
                                    '--right', str(right), '--config', str(root / 'config.json'),
                                    '--previous', str(root / 'first/复核工作簿.xlsx'), '--out', str(root / 'second')],
                                    capture_output=True, text=True)
            self.assertEqual(second.returncode, 0, second.stderr)
            self.assertEqual((root / 'config.json').read_bytes(), config)
            self.assertEqual(json.loads(second.stdout)['tasks']['carriedResolved'], 1)
            wb = openpyxl.load_workbook(root / 'second/复核工作簿.xlsx')
            self.assertEqual(wb['待处理']['E2'].value, '模拟确认：差额1已核实')
            self.assertEqual(wb['待处理']['F2'].value, '2')
            wb.close()
            summary = json.loads((root / 'second/摘要.json').read_text(encoding='utf-8'))
            self.assertEqual(summary['leftTotal'], '150.00')
            self.assertEqual(summary['leftRows'], 2)


if __name__ == '__main__':
    unittest.main()
