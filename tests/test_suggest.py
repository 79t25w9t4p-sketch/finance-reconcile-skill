import importlib.util
from pathlib import Path
import json
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
spec = importlib.util.spec_from_file_location('suggest', ROOT/'scripts/suggest.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

class Suggestions(unittest.TestCase):
    def workbook(self,path,sheets):
        wb=openpyxl.Workbook();wb.remove(wb.active)
        for name,rows in sheets:
            ws=wb.create_sheet(name)
            for row in rows:ws.append(row)
        wb.save(path);wb.close()

    def test_header_row_without_sheet_skips_short_instruction_pages(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'notes.xlsx'
            self.workbook(path,[('说明',[['说明']]),('明细',[['标题'],[],['编号','金额'],['A','10']])])
            info,location=module.inspect_input(path,header_row='3')
            self.assertEqual(location,{'sheet':'明细','header_row':'3'})
            self.assertEqual(info['rows'],1)
            with self.assertRaisesRegex(ValueError,'Header row outside sheet'):
                module.inspect_input(path,'说明','3')

    def test_hidden_only_candidate_requires_choice_but_explicit_sheet_is_honored(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'archive.xlsx'
            self.workbook(path,[('说明',[['说明']]),('历史',[['编号','金额'],['A','10']])])
            wb=openpyxl.load_workbook(path);wb['历史'].sheet_state='hidden';wb.save(path);wb.close()
            info,location=module.inspect_input(path)
            self.assertIsNone(location)
            self.assertEqual(info['candidates'][0]['state'],'hidden')
            result=subprocess.run([sys.executable,str(ROOT/'scripts/suggest.py'),'--left',str(path),'--right',str(path)],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            proposal=json.loads(result.stdout)
            self.assertEqual(proposal['status'],'needs_table_choice')
            self.assertIn('隐藏',proposal['questions'][0]['reason'])
            info,location=module.inspect_input(path,sheet='历史')
            self.assertEqual(location,{'sheet':'历史','header_row':'1'})
            wb=openpyxl.load_workbook(path);wb['历史'].sheet_state='veryHidden';wb.save(path);wb.close()
            info,location=module.inspect_input(path)
            self.assertIsNone(location)
            self.assertEqual(info['candidates'][0]['state'],'veryHidden')

    def test_finds_detail_sheet_and_nonfirst_header_without_printing_data(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'left.xlsx'
            self.workbook(path,[('使用说明',[['虚构说明']]),('销售明细',[['月末明细'],[],['送货单号','含税金额','未税金额'],['0001','1250.00','1106.19']])])
            result=subprocess.run([sys.executable,str(ROOT/'scripts/suggest.py'),'--left',str(path),'--right',str(path)],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            p=json.loads(result.stdout)
            self.assertEqual(p['status'],'needs_field_choice')
            self.assertEqual(p['mapping']['left_sheet'],'销售明细')
            self.assertEqual(p['mapping']['left_header_row'],'3')
            self.assertNotIn('left_amount',p['mapping'])
            self.assertEqual(p['inputOverview']['left']['rows'],1)
            self.assertNotIn('1250.00',result.stdout);self.assertNotIn('0001',result.stdout)

    def test_multiple_sheets_and_repeated_header_rows_require_choice(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ambiguous.xlsx'
            rows=[['编号','金额'],['A','10']]
            self.workbook(path,[('本月',rows),('上月',rows)])
            info,location=module.inspect_input(path)
            self.assertIsNone(location)
            self.assertEqual({x['sheet'] for x in info['candidates']},{'本月','上月'})
            result=subprocess.run([sys.executable,str(ROOT/'scripts/suggest.py'),'--left',str(path),'--right',str(path)],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertEqual(json.loads(result.stdout)['status'],'needs_table_choice')
            self.assertNotIn('mapping',json.loads(result.stdout))
            self.workbook(path,[('本月',rows+rows)])
            info,location=module.inspect_input(path)
            self.assertIsNone(location)
            self.assertEqual([x['headerRow'] for x in info['candidates']],[1,3])

    def test_unknown_headers_allow_explicit_selection_without_raw_preview(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'unknown.xlsx'
            self.workbook(path,[('导出',[['业务凭号','净额'],['私密内容','10']])])
            info,location=module.inspect_input(path)
            self.assertIsNone(location);self.assertEqual(info['candidates'],[])
            self.assertNotIn('私密内容',str(info))
            info,location=module.inspect_input(path,'导出','1')
            self.assertEqual(info['headers'],['业务凭号','净额'])
            self.assertEqual(location,{'sheet':'导出','header_row':'1'})
            self.workbook(path,[('--layout',[['业务凭号','净额'],['私密内容','10']])])
            info,location=module.inspect_input(path,'--layout','1')
            self.assertEqual(location['sheet'],'--layout')

    def test_scan_is_bounded_and_explicit_header_beyond_twenty_is_supported(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'late.xlsx'
            self.workbook(path,[('明细',[[str(i)] for i in range(20)]+[['编号','金额'],['A','10']])])
            info,location=module.inspect_input(path)
            self.assertIsNone(location);self.assertEqual(info['scanRows'],20)
            info,location=module.inspect_input(path,header_row='21')
            self.assertEqual(location,{'sheet':'明细','header_row':'21'})
            with self.assertRaisesRegex(ValueError,'正整数'):module.inspect_input(path,header_row='0')
            self.workbook(path,[(str(i),[['说明']]) for i in range(21)])
            with self.assertRaisesRegex(ValueError,'20个工作表'):module.inspect_input(path)
            long_headers=['编号','金额']+[str(i)+('长'*32000) for i in range(40)]
            self.workbook(path,[('过长',[long_headers,long_headers])])
            with self.assertRaisesRegex(ValueError,'候选表头超过2MB'):module.inspect_input(path)

    def test_scanning_keeps_formula_and_expanded_size_guards(self):
        import zipfile
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'formula.xlsx'
            self.workbook(path,[('明细',[['编号','金额'],['A','=1+1']])])
            with self.assertRaisesRegex(ValueError,'Formula'):module.inspect_input(path)
            path=Path(folder)/'large.xlsx'
            with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('large',b'0'*(40*1024*1024+1))
            with self.assertRaisesRegex(ValueError,'expanded size'):module.inspect_input(path)

    def test_source_change_during_layout_inspection_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'changed.xlsx'
            self.workbook(path,[('明细',[['编号','金额'],['A','10']])])
            original_run=subprocess.run
            def change_after_scan(command,**kwargs):
                result=original_run(command,**kwargs)
                if '--layout' in command:self.workbook(path,[('明细',[['编号','金额'],['A','20']])])
                return result
            with patch.object(module.subprocess,'run',side_effect=change_after_scan):
                with self.assertRaisesRegex(ValueError,'输入文件发生变化'):module.inspect_input(path)

    def test_delivery_numbers_and_tax_amounts_need_scope_confirmation(self):
        p=module.propose(['送货单号','含税金额'],['供方送货单号','含税金额'])
        self.assertEqual(p['status'],'ready_for_confirmation')
        self.assertEqual(json.loads(p['mapping']['left_keys']),['送货单号'])
        self.assertEqual(json.loads(p['mapping']['right_keys']),['供方送货单号'])
        self.assertIn('两表期间、币种、金额含义和明细粒度一致',p['confirm'])
        ambiguous=module.propose(['送货单号','含税金额','未税金额'],['供方送货单号','含税金额'])
        self.assertEqual(ambiguous['status'],'needs_field_choice')
        self.assertNotIn('left_amount',ambiguous['mapping'])
    def test_different_headers_map_same_roles(self):
        p=module.propose(['客户','订单号','应收金额'],['客户名称','单据编号','回款金额'])
        self.assertEqual(p['status'],'ready_for_confirmation')
        self.assertEqual(json.loads(p['mapping']['left_keys']),['客户','订单号'])
        self.assertEqual(p['mapping']['right_amount'],'回款金额')
    def test_ambiguous_amounts_and_ids_require_choice(self):
        p=module.propose(['订单号','发票号','借方金额','贷方金额'],['单号','金额'])
        self.assertEqual(p['status'],'needs_field_choice')
        self.assertNotIn('left_amount',p['mapping']);self.assertNotIn('left_keys',p['mapping'])
    def test_codes_do_not_match_names(self):
        for left,right in [(['客户编码','编号','金额'],['客户名称','编号','金额']),(['公司','编号','金额'],['编号','金额'])]:
            self.assertEqual(module.propose(left,right)['status'],'needs_field_choice')
    def test_company_and_shared_customer_code_are_kept(self):
        headers=['公司','客户编号','客户名称','单号','金额']
        p=module.propose(headers,headers)
        self.assertEqual(json.loads(p['mapping']['left_keys']),['公司','客户编号','单号'])
    def test_unknown_headers_are_not_fuzzy_guessed(self):
        p=module.propose(['业务备注','未税净额'],['说明','含税总计'])
        self.assertEqual(p['status'],'needs_field_choice')
        self.assertEqual(p['mapping'],{})
    def test_cli_is_read_only_and_does_not_print_detail(self):
        result=subprocess.run([sys.executable,str(ROOT/'scripts/suggest.py'),'--left',str(ROOT/'examples/左表.csv'),'--right',str(ROOT/'examples/右表.csv')],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        p=json.loads(result.stdout);self.assertEqual(p['inputOverview']['left']['rows'],7)
        self.assertNotIn('AR001',result.stdout)
        self.assertEqual(json.loads(p['mapping']['left_keys']),['客户','编号'])

if __name__=='__main__': unittest.main()
