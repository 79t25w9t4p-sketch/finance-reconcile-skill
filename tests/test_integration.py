# coding: utf-8
import unittest, tempfile, subprocess, json
from pathlib import Path
import sys
import hashlib
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
class Integration(unittest.TestCase):
    def run_cli(self,*args):
        return subprocess.run([sys.executable,str(ROOT/'scripts/reconcile.py'),*map(str,args)],capture_output=True,text=True)
    def workbook(self,path,formula=False,multiple=False):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='明细'
        ws.append(['编号','金额']);ws.append(['001','=1+1' if formula else 0.30])
        if multiple: wb.create_sheet('其他')
        wb.save(path)
    def test_csv_report_and_overwrite_refused(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'report'
            args=['--single-key-confirmed','--left',ROOT/'examples/左表.csv','--right',ROOT/'examples/右表.csv','--left-key','编号','--left-amount','金额','--right-key','单据号','--right-amount','余额','--out',out]
            r=self.run_cli(*args);self.assertEqual(r.returncode,0,r.stderr)
            summary=json.loads(r.stdout);self.assertEqual(summary['counts']['重复编号'],1)
            report=json.loads((out/'核对记录.json').read_text());self.assertEqual(next(x for x in report['result'] if x['id']=='AR003')['delta'],'10')
            manifest=json.loads((out/'运行清单.json').read_text());self.assertTrue(manifest['coverage']['left']['everyRecordExactlyOnce']);self.assertEqual(manifest['coverage']['left']['reportedRows'],7);self.assertEqual(manifest['inputs']['left']['sha256'],hashlib.sha256((ROOT/'examples/左表.csv').read_bytes()).hexdigest());self.assertEqual(manifest['outputs'][0]['sha256'],hashlib.sha256((out/'差异表.csv').read_bytes()).hexdigest())
            before=(out/'差异表.csv').read_bytes();self.assertNotEqual(self.run_cli(*args).returncode,0);self.assertEqual(before,(out/'差异表.csv').read_bytes())
    def test_xlsx_exact_text_identifier(self):
        with tempfile.TemporaryDirectory() as d:
            left=Path(d)/'left.xlsx';right=Path(d)/'right.csv';self.workbook(left);right.write_text('编号,金额\n001,0.29',encoding='utf8')
            r=self.run_cli('--left',left,'--right',right,'--left-key','编号','--left-amount','金额','--right-key','编号','--right-amount','金额','--out',Path(d)/'result')
            self.assertEqual(r.returncode,0,r.stderr)
            report=json.loads((Path(d)/'result/核对记录.json').read_text());self.assertEqual(report['result'][0]['delta'],'1');self.assertEqual(report['result'][0]['id'],'001')
    def test_xlsx_header_offset_and_original_rows(self):
        with tempfile.TemporaryDirectory() as d:
            left=Path(d)/'offset.xlsx';right=Path(d)/'right.csv'
            wb=openpyxl.Workbook();ws=wb.active;ws.append(['导出说明']);ws.append([]);ws.append(['编号','金额']);ws.append(['001',0.30]);ws.append([]);ws.append(['002',2]);wb.save(left)
            right.write_text('编号,金额\n001,0.29\n002,2',encoding='utf8')
            r=self.run_cli('--left',left,'--right',right,'--left-header-row','3','--left-key','编号','--left-amount','金额','--right-key','编号','--right-amount','金额','--out',Path(d)/'result')
            self.assertEqual(r.returncode,0,r.stderr)
            report=json.loads((Path(d)/'result/核对记录.json').read_text());self.assertEqual([x['left'][0]['line'] for x in report['result']],[4,6])
            manifest=json.loads((Path(d)/'result/运行清单.json').read_text());self.assertEqual(manifest['inputs']['left']['headerRow'],3);self.assertEqual(manifest['inputs']['left']['sheet'],'Sheet')
    def test_help_requires_no_input(self):
        r=self.run_cli('--help');self.assertEqual(r.returncode,0);self.assertRegex(r.stdout,r'财务明细核对 \d+\.\d+\.\d+')
    def test_composite_keys_prevent_cross_customer_match(self):
        with tempfile.TemporaryDirectory() as d:
            left=Path(d)/'left.csv';right=Path(d)/'right.csv'
            left.write_text('客户,编号,金额\n甲,001,100',encoding='utf8');right.write_text('客户名称,单号,金额\n乙,001,100',encoding='utf8')
            base=['--left',left,'--right',right,'--left-amount','金额','--right-amount','金额']
            blocked=self.run_cli(*base,'--left-key','编号','--right-key','单号','--out',Path(d)/'blocked');self.assertNotEqual(blocked.returncode,0);self.assertFalse((Path(d)/'blocked').exists())
            r=self.run_cli(*base,'--left-keys','["客户","编号"]','--right-keys','["客户名称","单号"]','--out',Path(d)/'result');self.assertEqual(r.returncode,0,r.stderr);self.assertEqual(json.loads(r.stdout)['counts'],{'仅左表':1,'仅右表':1})
    def test_formula_refused(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'formula.xlsx';self.workbook(p,formula=True)
            r=self.run_cli('--left',p,'--right',p,'--inspect');self.assertNotEqual(r.returncode,0);self.assertIn('Formula',r.stderr)
    def test_multisheet_requires_selection(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'multi.xlsx';self.workbook(p,multiple=True)
            r=self.run_cli('--left',p,'--right',p,'--inspect');self.assertNotEqual(r.returncode,0)
            r=self.run_cli('--left',p,'--right',p,'--left-sheet','明细','--right-sheet','明细','--inspect');self.assertEqual(r.returncode,0,r.stderr)
if __name__=='__main__': unittest.main()
