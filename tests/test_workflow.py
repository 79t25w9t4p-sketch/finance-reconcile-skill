import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import workflow as workflow_module

class Workflow(unittest.TestCase):
    def test_both_workbooks_show_actual_source_locations_as_literal_text(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'=left.xlsx';right=d/'right.xlsx';config=d/'config.json'
            for path,name,header,amount in [(left,'=销售明细',3,'10'),(right,'采购明细',1,'9')]:
                wb=openpyxl.Workbook();ws=wb.active;ws.title=name
                for _ in range(header-1):ws.append(['说明'])
                ws.append(['编号','金额']);ws.append(['001',amount]);wb.save(path);wb.close()
            base=[sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right)]
            result=subprocess.run(base+['--left-sheet','=销售明细','--left-header-row','3',
                                       '--left-keys','["编号"]','--right-keys','["编号"]',
                                       '--left-amount','金额','--right-amount','金额',
                                       '--save-config',str(config),'--out',str(d/'first')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            saved=config.read_bytes()
            self.assertNotIn('right_sheet',json.loads(saved)['mapping'])
            for file,sheet in [('复核工作簿.xlsx','汇总'),('跟进清单.xlsx','使用说明')]:
                wb=openpyxl.load_workbook(d/'first'/file)
                info=dict(wb[sheet].values)
                self.assertEqual(info['左表文件'],'=left.xlsx')
                self.assertEqual(info['左表工作表'],'=销售明细');self.assertEqual(str(info['左表表头行']),'3')
                self.assertEqual(info['右表工作表'],'采购明细');self.assertEqual(str(info['右表表头行']),'1')
                self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
                wb.close()
            previous=d/'first/复核工作簿.xlsx'
            wb=openpyxl.load_workbook(previous);wb['待处理']['D2']='已解决';wb['待处理']['E2']='核实过';wb.save(previous);wb.close()
            result=subprocess.run(base+['--config',str(config),'--previous',str(previous),'--out',str(d/'second')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertEqual(json.loads(result.stdout)['tasks']['carriedResolved'],1)
            self.assertEqual(config.read_bytes(),saved)

    def test_source_change_between_inspection_and_report_leaves_no_output(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.csv';right=d/'right.csv'
            left.write_text('编号,金额\nA,10',encoding='utf-8');right.write_text('编号,金额\nA,9',encoding='utf-8')
            original_run=subprocess.run
            def change_then_run(command,**kwargs):
                if '--out' in command:left.write_text('编号,金额\nA,20',encoding='utf-8')
                return original_run(command,**kwargs)
            argv=['workflow.py','--left',str(left),'--right',str(right),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--save-config',str(d/'config.json'),'--out',str(d/'out')]
            with patch.object(sys,'argv',argv),patch.object(workflow_module.subprocess,'run',side_effect=change_then_run):
                with self.assertRaisesRegex(ValueError,'输入文件发生变化'):workflow_module.main()
            self.assertFalse((d/'out').exists());self.assertFalse((d/'config.json').exists())

    def test_followup_literal_ids_and_empty_pending(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.csv';right=d/'right.csv';config=d/'config.json'
            left.write_text('编号,金额\n=单据,10',encoding='utf8')
            right.write_text('编号,金额\n=单据,9',encoding='utf8')
            base=[sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right)]
            result=subprocess.run(base+['--save-config',str(config),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--out',str(d/'first')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            followup=openpyxl.load_workbook(d/'first/跟进清单.xlsx')
            self.assertEqual(followup['待核实事项']['A2'].value,'=单据')
            self.assertEqual(followup['待核实事项']['A2'].data_type,'s')
            right.write_text('编号,金额\n=单据,10',encoding='utf8')
            result=subprocess.run(base+['--config',str(config),'--out',str(d/'second')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            followup=openpyxl.load_workbook(d/'second/跟进清单.xlsx')
            self.assertEqual(followup['待核实事项'].max_row,1)
            self.assertFalse(any(c.data_type=='f' for ws in followup for row in ws for c in row))
    def test_reuse_review_changes_history_and_schema(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.csv';right=d/'right.csv';config=d/'配置.json'
            left.write_text('编号,金额,说明\nA,10,说明\nB,20,说明\nC,30,说明\nD,40,说明',encoding='utf8')
            right.write_text('编号,金额,说明\nA,9,说明\nB,19,说明\nC,29,说明\nD,39,说明',encoding='utf8')
            def run(out,*extra):
                return subprocess.run([sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right),'--out',str(d/out),*map(str,extra)],capture_output=True,text=True)
            r=run('first','--save-config',config,'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额');self.assertEqual(r.returncode,0,r.stderr)
            saved=json.loads(config.read_text());self.assertNotIn(str(d),config.read_text())
            previous=d/'first/复核工作簿.xlsx';wb=openpyxl.load_workbook(previous);ws=wb['待处理']
            for row in range(2,6):ws.cell(row,4,'已解决');ws.cell(row,5,'=原始意见')
            # Literal text, not an Excel formula.
            for row in range(2,6):ws.cell(row,5).data_type='s'
            wb.save(previous)
            left.write_text('编号,金额,说明\nB,21,说明\nA,10,说明\nC,30,改变了说明\nE,50,新增',encoding='utf8')
            right.write_text('编号,金额,说明\nE,49,新增\nC,29,说明\nA,9,说明\nB,19,说明',encoding='utf8')
            r=run('second','--config',config,'--previous',previous);self.assertEqual(r.returncode,0,r.stderr)
            wb=openpyxl.load_workbook(d/'second/复核工作簿.xlsx');rows={row[0].value:row for row in list(wb['待处理'])[1:]}
            self.assertEqual(rows['A'][2].value,'未变化');self.assertEqual(rows['A'][3].value,'已解决');self.assertEqual(rows['A'][4].value,'=原始意见');self.assertEqual(rows['A'][4].data_type,'s')
            self.assertEqual(rows['B'][3].value,'需重新复核');self.assertEqual(rows['C'][3].value,'需重新复核');self.assertEqual(rows['E'][2].value,'新增')
            self.assertEqual(wb['汇总']['B2'].value,3);self.assertEqual(wb['汇总']['B3'].value,2);self.assertEqual(wb['汇总']['B4'].value,1);self.assertEqual(wb['汇总']['B5'].value,1);self.assertEqual(wb['待处理']['A2'].value,'B');
            self.assertEqual(wb['历史记录']['A2'].value,'D');self.assertEqual(wb['历史记录']['E2'].value,'=原始意见')
            focus = wb['本次重点']
            self.assertEqual({r[0] for r in list(focus.values)[1:]}, {'B','C','E'})
            self.assertTrue(focus.protection.sheet)
            for row in list(focus)[1:]:
                destination = row[8].hyperlink.location
                cell = destination.split('!')[1]
                self.assertEqual(wb['待处理'][cell].row, next(r[0].row for r in list(wb['待处理'])[1:] if r[0].value == row[0].value))
            self.assertTrue(wb['待处理'].auto_filter.ref);self.assertEqual(wb['待处理'].freeze_panes,'F2');self.assertEqual(len(wb['待处理'].data_validations.dataValidation),1)
            followup=openpyxl.load_workbook(d/'second/跟进清单.xlsx')
            self.assertEqual(followup.sheetnames,['使用说明','待核实事项'])
            self.assertEqual(followup['待核实事项'].max_row,4)
            followup_rows=list(followup['待核实事项'].values)
            self.assertEqual({r[0] for r in followup_rows[1:]},{'B','C','E'})
            self.assertNotIn('=原始意见',str([list(ws.values) for ws in followup]))
            self.assertTrue(all(r[2] for r in followup_rows[1:]))
            manifest=json.loads((d/'second/运行清单.json').read_text())
            self.assertIn('跟进清单.xlsx',[x['file'] for x in manifest['outputs']])
            self.assertEqual(manifest['reviewTasks'], {'focus':3,'pending':3,'firstPending':0,'newPending':1,'recheck':2,'ongoing':0,'carriedResolved':1,'absent':1})
            self.assertIn('本次待办摘要.txt',[x['file'] for x in manifest['outputs']])
            right.write_text('编号,余额,说明\nA,9,说明',encoding='utf8')
            r=run('invalid','--config',config);self.assertNotEqual(r.returncode,0);self.assertIn('表头',r.stderr);self.assertFalse((d/'invalid').exists());self.assertEqual(saved,json.loads(config.read_text()))

    def test_invalid_previous_leaves_no_partial_output(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);(d/'left.csv').write_text('编号,金额\nA,10');(d/'right.csv').write_text('编号,金额\nA,9');(d/'bad.xlsx').write_text('not a workbook')
            r=subprocess.run([sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(d/'left.csv'),'--right',str(d/'right.csv'),'--out',str(d/'out'),'--save-config',str(d/'config.json'),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--previous',str(d/'bad.xlsx')],capture_output=True,text=True)
            self.assertNotEqual(r.returncode,0);self.assertFalse((d/'out').exists());self.assertFalse((d/'config.json').exists())

    def test_unchanged_pending_is_not_lost_when_focus_is_empty(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.csv';right=d/'right.csv';config=d/'config.json'
            left.write_text('编号,金额\nA,10\nB,20',encoding='utf-8')
            right.write_text('编号,金额\nA,9\nB,19',encoding='utf-8')
            base=[sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right)]
            result=subprocess.run(base+['--save-config',str(config),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--out',str(d/'first')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            previous=d/'first/复核工作簿.xlsx';wb=openpyxl.load_workbook(previous)
            wb['待处理']['D2']='已解决';wb['待处理']['E2']='已核实';wb.save(previous)
            result=subprocess.run(base+['--config',str(config),'--previous',str(previous),'--out',str(d/'second')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            tasks=json.loads(result.stdout)['tasks']
            self.assertEqual(tasks['focus'],0);self.assertEqual(tasks['ongoing'],1);self.assertEqual(tasks['carriedResolved'],1)
            wb=openpyxl.load_workbook(d/'second/复核工作簿.xlsx')
            self.assertEqual(wb['本次重点'].max_row,1)
            followup=openpyxl.load_workbook(d/'second/跟进清单.xlsx')
            self.assertEqual(followup['待核实事项']['A2'].value,'B');self.assertEqual(followup['待核实事项'].max_row,2)

    def test_runs_without_any_executable_on_path_and_preserves_huge_delta(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.xlsx';right=d/'right.xlsx'
            huge='9'*60+'.01'
            for path,amount in [(left,huge),(right,'0')]:
                wb=openpyxl.Workbook();ws=wb.active;ws.append(['编号','金额']);ws.append(['001',amount]);wb.save(path)
            result=subprocess.run([sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right),'--left-keys','["编号"]','--right-keys','["编号"]','--left-amount','金额','--right-amount','金额','--save-config',str(d/'config.json'),'--out',str(d/'out')],capture_output=True,text=True,env={**os.environ,'PATH':'','PYTHONUTF8':'1'})
            self.assertEqual(result.returncode,0,result.stderr)
            wb=openpyxl.load_workbook(d/'out/复核工作簿.xlsx')
            self.assertEqual(wb['待处理']['J2'].value,huge)
            self.assertEqual(json.loads(result.stdout)['tasks']['firstPending'],1)

    def test_v07_review_and_profile_can_be_reused(self):
        with tempfile.TemporaryDirectory() as folder:
            d=Path(folder);left=d/'left.csv';right=d/'right.csv';config=d/'config.json'
            left.write_text('编号,金额\nA,100\nB,200\nD,400',encoding='utf-8')
            right.write_text('编号,金额\nA,99\nB,199\nD,399',encoding='utf-8')
            profile={'version':1,'mapping':{'left_keys':'["编号"]','right_keys':'["编号"]','left_amount':'金额','right_amount':'金额'},'headers':{'left':['编号','金额'],'right':['编号','金额']}}
            config.write_text(json.dumps(profile,ensure_ascii=False),encoding='utf-8')
            previous=ROOT/'examples/首次复核示例.xlsx'
            result=subprocess.run([sys.executable,str(ROOT/'scripts/workflow.py'),'--left',str(left),'--right',str(right),'--config',str(config),'--previous',str(previous),'--out',str(d/'out')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            wb=openpyxl.load_workbook(d/'out/复核工作簿.xlsx')
            rows={r[0]:r for r in list(wb['待处理'].values)[1:]}
            self.assertEqual(rows['A'][3],'已解决');self.assertEqual(rows['B'][3],'已解决')
            self.assertEqual(rows['D'][2],'新增')

if __name__ == '__main__': unittest.main()
