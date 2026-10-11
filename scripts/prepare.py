#!/usr/bin/env python3
"""Prepare a value-only worksheet after explicit non-detail-row confirmation."""
import argparse
import json
from pathlib import Path
import re
import sys
import tempfile
import zipfile

import openpyxl
from reconcile import fingerprint, load_table
from suggest import data_warnings
from workflow import text_cell


def preserve_carriage_returns(path):
    """Escape literal CRs so XML readers preserve values with either writer backend."""
    sheet_xml = 'xl/worksheets/sheet1.xml'
    fixed = path.with_name('.literal-newlines.xlsx')
    with zipfile.ZipFile(path) as original:
        if b'\r' not in original.read(sheet_xml):
            return
        with zipfile.ZipFile(fixed, 'w') as replacement:
            for member in original.infolist():
                value = original.read(member.filename)
                if member.filename == sheet_xml:
                    value = value.replace(b'\r', b'&#13;')
                replacement.writestr(member, value)
    fixed.replace(path)


def prepare(file, sheet=None, header_row='1', *, expected_sha256=None, exclude_rows=None,
            exclude_kinds=None, confirmed=False, reason=None, out=None):
    source = Path(file)
    table, metadata = load_table(source, sheet, header_row)
    warnings = data_warnings(table, position_limit=None)
    if out is None:
        return {'source': metadata, 'dataWarnings': warnings,
                'positionKind': 'xlsx_row' if metadata['sheet'] is not None else 'parsed_record',
                'note': '只读检查。位置已完整列出；CSV/TSV位置是非空解析记录序号，不是物理文本行号。'}

    if not confirmed or not isinstance(reason, str) or not reason.strip():
        raise ValueError('须明确确认非明细行并记录排除理由，不能自行删行')
    if not isinstance(expected_sha256, str) or not re.fullmatch(r'[0-9a-f]{64}', expected_sha256):
        raise ValueError('须提供本次预检的SHA-256')
    if expected_sha256 != metadata['sha256']:
        raise ValueError('输入已变化，请重新预检和确认，不能沿用旧排除位置')
    if (exclude_rows is None) == (exclude_kinds is None):
        raise ValueError('明确排除位置或已确认的整类行，二选一')
    candidates = {position: warning['kind'] for warning in warnings for position in warning['rows']}
    if exclude_kinds is not None:
        if not exclude_kinds or not set(exclude_kinds).issubset({'summary_rows', 'repeated_headers'}):
            raise ValueError('排除类别无效')
        selected = sorted(position for position, kind in candidates.items() if kind in exclude_kinds)
        if any(not any(w['kind'] == kind for w in warnings) for kind in exclude_kinds):
            raise ValueError('本次输入不存在指定的疑似行类别')
    else:
        if not isinstance(exclude_rows, list) or not exclude_rows or any(type(n) is not int for n in exclude_rows):
            raise ValueError('排除位置必须为非空整数数组')
        if len(set(exclude_rows)) != len(exclude_rows):
            raise ValueError('排除位置重复')
        selected = sorted(exclude_rows)
        if any(n not in candidates for n in selected):
            raise ValueError('排除位置不是本次预检的疑似汇总行或重复表头，不能删除普通明细或表头')
    if len(selected) == len(table['rows']):
        raise ValueError('排除后没有明细，停止制作副本')

    destination = Path(out)
    if destination.exists():
        raise ValueError('输出目录已存在，不能覆盖原文件或已有结果')
    positions = list(table.get('sourceLines', range(2, len(table['rows']) + 2)))
    excluded = set(selected)
    retained = [(n, row) for n, row in zip(positions, table['rows']) if n not in excluded]
    copy_sheet = metadata['sheet'] or '核对明细'
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.prepare-', dir=destination.parent) as staging:
        stage = Path(staging)
        copy_path = stage / '核对副本.xlsx'
        workbook = openpyxl.Workbook()
        try:
            worksheet = workbook.active
            worksheet.title = copy_sheet
            for col, value in enumerate(table['headers'], 1):
                text_cell(worksheet.cell(metadata['headerRow'], col), value)
            for row_number, row in retained:
                for col, value in enumerate(row, 1):
                    text_cell(worksheet.cell(row_number, col), value)
            worksheet.freeze_panes = 'A' + str(metadata['headerRow'] + 1)
            workbook.save(copy_path)
        finally:
            workbook.close()
        preserve_carriage_returns(copy_path)
        # Re-read before publishing: reject any lossy conversion or unreadable copy.
        check, copy_metadata = load_table(copy_path, copy_sheet, str(metadata['headerRow']))
        if (check['headers'] != table['headers'] or check['rows'] != [row for _, row in retained]
                or check['sourceLines'] != [n for n, _ in retained]):
            raise ValueError('副本内容或来源位置不一致，停止输出')
        audit = {'format': 'finance-preparation-v1', 'source': metadata,
                 'copy': copy_metadata, 'positionKind': 'xlsx_row' if metadata['sheet'] is not None else 'parsed_record',
                 'excludedRows': [{'position': n, 'kind': candidates[n]} for n in selected],
                 'retainedPositions': [n for n, _ in retained], 'reason': reason.strip(),
                 'note': '仅复制所选明细页的读取值与列名，不复制格式或其他工作表。原文件不改。副本行号对应原XLSX行号或CSV/TSV非空解析记录序号；不是CSV物理文本行号。'}
        (stage / '整理说明.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
        if fingerprint(source, 10 * 1024 * 1024) != expected_sha256:
            raise ValueError('整理过程中输入发生变化，停止输出，请重新预检和确认')
        # Never replace a directory that appeared during preparation.
        destination.mkdir()
        try:
            for path in stage.iterdir():
                path.rename(destination / path.name)
        except Exception:
            for name in ['核对副本.xlsx', '整理说明.json']:
                (destination / name).unlink(missing_ok=True)
            destination.rmdir()
            raise
    return {'source': metadata, 'copy': copy_metadata, 'out': str(destination),
            'excludedCount': len(selected), 'remainingWarnings': data_warnings(check),
            'note': '请对核对副本重新预检后运行工作流；保留原文件与整理说明，来源行指向副本。'}


def main():
    parser = argparse.ArgumentParser(description='只读列出疑似行，或确认后制作新的核对副本')
    parser.add_argument('--input', required=True)
    parser.add_argument('--sheet')
    parser.add_argument('--header-row', default='1')
    parser.add_argument('--inspect', action='store_true')
    parser.add_argument('--expect-sha256')
    parser.add_argument('--exclude-rows', help='本次已确认的来源位置JSON数组')
    parser.add_argument('--exclude-kinds', nargs='+', choices=['summary_rows', 'repeated_headers'],
                        help='仅在整类疑似行均已确认非明细时使用，包含全部位置')
    parser.add_argument('--confirmed', action='store_true')
    parser.add_argument('--reason')
    parser.add_argument('--out', help='尚不存在的副本目录')
    args = parser.parse_args()
    if args.inspect:
        if any([args.out, args.expect_sha256, args.exclude_rows, args.exclude_kinds, args.confirmed, args.reason]):
            raise ValueError('只读检查不能同时传入整理参数')
    elif not args.out:
        raise ValueError('使用 --inspect 只读检查，或用 --out 指定新的副本目录')
    result = prepare(args.input, args.sheet, args.header_row, expected_sha256=args.expect_sha256,
                     exclude_rows=json.loads(args.exclude_rows) if args.exclude_rows is not None else None,
                     exclude_kinds=args.exclude_kinds, confirmed=args.confirmed, reason=args.reason, out=args.out)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print('副本准备失败：' + str(error), file=sys.stderr)
        sys.exit(1)
