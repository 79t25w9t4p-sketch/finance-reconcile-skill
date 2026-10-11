#!/usr/bin/env python3
"""Repeatable local reconciliation and review workbook workflow."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile
from core import fmt
from reconcile import fingerprint as file_hash

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parent
VERSION = '0.8.7'
HEADERS = ['编号 / 联合键', '核对结果', '本次变化', '处理状态', '处理意见', '左表来源行', '右表来源行', '左表原金额', '右表原金额', '左减右差额', '记录标识', '内容指纹']
STATES = {'待处理', '处理中', '已解决', '需重新复核', '无需处理'}


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def read_previous(path, profile_hash):
    if not path:
        return {}, []
    path = Path(path)
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('上次工作簿超过10MB')
    with zipfile.ZipFile(path) as z:
        if sum(x.file_size for x in z.infolist()) > 40 * 1024 * 1024:
            raise ValueError('上次工作簿解压后超过40MB')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if wb['_meta'].max_row != 2 or wb['_meta'].max_column != 2:
            raise ValueError('工作簿标识结构无效')
        meta = dict(wb['_meta'].iter_rows(values_only=True))
        if meta.get('format') != 'finance-review-v1' or meta.get('profile') != profile_hash:
            raise ValueError('上次工作簿的配置不一致；不能自动沿用意见')
        previous = {}
        history = []
        for name in ['待处理', '已匹配', '历史记录']:
            ws = wb[name]
            if ws.max_row > 10001 or ws.max_column > 12:
                raise ValueError('上次工作簿结构或记录数量异常')
            if list(next(ws.values)) != HEADERS:
                raise ValueError('请保留复核工作表的原表头')
            for cells in ws.iter_rows(min_row=2):
                if any(c.data_type == 'f' for c in cells):
                    raise ValueError('处理意见中请填写文本，不能使用公式')
                row = [c.value if c.value is not None else '' for c in cells]
                if not any(row):
                    continue
                if row[3] not in STATES:
                    raise ValueError('处理状态无效，请使用下拉选项')
                if name == '历史记录':
                    history.append(row)
                else:
                    if not row[10] or row[10] in previous:
                        raise ValueError('上次工作簿存在缺失或重复记录标识')
                    previous[row[10]] = row
        return previous, history
    finally:
        wb.close()


def text_cell(cell, value):
    # Treat all strings as literal text, including =, +, -, and @ prefixes.
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError('单元格内容过长，请按较小范围核对')
        cell.value = value
        cell.data_type = 's'
    else:
        cell.value = value


def make_sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    for row_number, row in enumerate([HEADERS] + rows, 1):
        for col, value in enumerate(row, 1):
            text_cell(ws.cell(row_number, col), value)
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = ws.dimensions
    widths = [34, 18, 20, 18, 48, 18, 18, 22, 22, 20, 20, 20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in row[3:5]:
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
    ws.column_dimensions['K'].hidden = True
    ws.column_dimensions['L'].hidden = True
    if rows:
        dv = DataValidation(type='list', formula1='"待处理,处理中,已解决,需重新复核,无需处理"')
        dv.showErrorMessage = True
        dv.error = '请从下拉选项中选择处理状态'
        ws.add_data_validation(dv)
        dv.add(f'D2:D{len(rows)+1}')
    return ws


def source_locations(inputs):
    rows=[]
    for side,label in [('left','左表'),('right','右表')]:
        source=inputs[side]
        rows.extend([[label+'文件',source['file']],
                     [label+'工作表',source['sheet'] if source['sheet'] is not None else 'CSV/TSV（无工作表）'],
                     [label+'表头行',source['headerRow']]])
    return rows


def save_followup(rows, profile, inputs, dest):
    """Export unresolved items without internal review notes or hidden identifiers."""
    advice = {
        '金额差异': '请核对两边计价、含税口径及优惠/退货调整，并回查原单据。',
        '仅左表': '请确认右表是否遗漏、是否跨期，或业务编号是否不同。',
        '仅右表': '请确认左表是否遗漏、是否跨期，或业务编号是否不同。',
        '重复编号': '请核实重复是重复登记还是一单多行；确认前不能直接删行或合并。',
        '无效金额': '请补齐或修正原表金额；空值不等于零。',
        '空编号': '请补齐业务编号；不能仅凭相同金额认定为同一笔。',
        '空编号及无效金额': '请补齐业务编号及金额后重新核对。',
        '一致': '记录内容变化后仍需复核；当前指定金额列相同不代表全部业务信息相同。',
    }
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = '使用说明'
    mapping = profile['mapping']
    info_rows = [
        ['待核实跟进清单', VERSION],
        ['用途', '列出生成时仍待处理的事项，供人工核实；不判定哪方错账。'],
        *source_locations(inputs),
        ['左表比较金额列', mapping['left_amount']], ['右表比较金额列', mapping['right_amount']],
        ['编号对应', '左：'+mapping['left_keys']+'；右：'+mapping['right_keys']],
        ['差额方向', '左表金额减右表金额；重复或异常不计算差额。'],
        ['分享前', '检查客户、期间、币种、含税口径和收件范围；此文件仍含业务编号与金额。'],
        ['状态口径', '只含生成时待处理、处理中、需重新复核项；不含已解决、无需处理和历史条目。'],
        ['反馈填写', '核实反馈供填写；需要沿用意见时，将结果填写回内部复核工作簿的黄色列。'],
    ]
    for number, row in enumerate(info_rows, 1):
        for col, value in enumerate(row, 1):
            text_cell(info.cell(number, col), value)
            info.cell(number, col).alignment = Alignment(wrap_text=True, vertical='top')
        info.row_dimensions[number].height = 32
    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 90
    info['A1'].font = Font(size=16, bold=True, color='17365D')
    ws = wb.create_sheet('待核实事项')
    headers = ['核对编号', '核对结果', '建议核实事项', '左表原金额', '右表原金额',
               '左减右差额', '左表来源行', '右表来源行', '核实反馈']
    pending = [r for r in rows if r[3] in {'待处理', '处理中', '需重新复核'}]
    table = [headers] + [[r[0], r[1], advice[r[1]], r[7], r[8], r[9], r[5], r[6], ''] for r in pending]
    for number, row in enumerate(table, 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(number, col)
            text_cell(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if number == 1:
                cell.fill = PatternFill('solid', fgColor='17365D')
                cell.font = Font(bold=True, color='FFFFFF')
            elif col == 9:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
        ws.row_dimensions[number].height = 48 if number > 1 else 30
    for col, width in enumerate([34, 18, 58, 20, 20, 20, 18, 18, 44], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(dest)


def make_focus(wb, rows, review_rows):
    """Read-only view; links point to the single authoritative editable record."""
    ws = wb.create_sheet('本次重点')
    headers = ['编号 / 联合键', '核对结果', '本次变化', '处理状态', '上次处理意见',
               '左表原金额', '右表原金额', '左减右差额', '处理入口']
    for col, value in enumerate(headers, 1):
        text_cell(ws.cell(1, col), value)
        ws.cell(1, col).fill = PatternFill('solid', fgColor='17365D')
        ws.cell(1, col).font = Font(bold=True, color='FFFFFF')
    for number, row in enumerate(rows, 2):
        values = [row[0], row[1], row[2], row[3], row[4], row[7], row[8], row[9], '填写状态与意见']
        for col, value in enumerate(values, 1):
            text_cell(ws.cell(number, col), value)
            ws.cell(number, col).alignment = Alignment(wrap_text=True, vertical='top')
        link = ws.cell(number, 9)
        # Internal location only: https://openpyxl.readthedocs.io/en/stable/api/openpyxl.worksheet.hyperlink.html
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'待处理'!D{review_rows[row[10]]}", display=link.value)
        link.style = 'Hyperlink'
        ws.row_dimensions[number].height = 38
    for col, width in enumerate([34, 18, 20, 18, 48, 22, 22, 20, 24], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = ws.dimensions
    # Prevent accidental view edits; this is not encryption or access control.
    # https://openpyxl.readthedocs.io/en/stable/protection.html
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    return ws


def build_workbook(report, summary, profile, previous_path, dest, inputs):
    profile_hash = digest(profile)
    previous, history = read_previous(previous_path, profile_hash)
    rows = []
    seen = set()
    changes = {}
    for index, item in enumerate(report['result']):
        ambiguous = '空编号' in item['status'] or item['status'] == '重复编号'
        identity = digest({'id': item['id'], 'blankIndex': index if '空编号' in item['status'] else None})
        contents = {'status': item['status'], 'delta': item['delta']}
        for side in ['left', 'right']:
            contents[side] = sorted([x.get('values', [x['id'], x['raw']]) for x in item[side]])
        fingerprint = digest(contents)
        old = previous.get(identity) if '空编号' not in item['status'] else None
        if '空编号' not in item['status']:
            seen.add(identity)
        note = old[4] if old else ''
        if old and old[11] == fingerprint and not ambiguous:
            change, state = '未变化', old[3]
        elif old:
            change, state = '需重新复核', '需重新复核'
        else:
            change = '新增' if previous_path else '首次核对'
            state = '无需处理' if item['status'] == '一致' else '待处理'
        changes[change] = changes.get(change, 0) + 1
        row = [item['id'], item['status'], change, state, note,
               ' / '.join(str(x['line']) for x in item['left']),
               ' / '.join(str(x['line']) for x in item['right']),
               ' / '.join(x['raw'] for x in item['left']),
               ' / '.join(x['raw'] for x in item['right']),
               '待核实' if item['delta'] is None else fmt(int(item['delta']), grouping=False),
               identity, fingerprint]
        rows.append(row)
    absent = 0
    for identity, old in previous.items():
        if identity not in seen:
            old[2] = '本次未出现'
            history.append(old)
            absent += 1
    if len(history) > 10000:
        raise ValueError('历史记录超过10000条，请另存归档后开始新周期')
    priority = {'需重新复核': 0, '待处理': 1, '处理中': 2, '已解决': 3, '无需处理': 4}
    rows.sort(key=lambda r: (priority[r[3]], 0 if r[2] == '新增' else 1, r[0]))
    attention = sum(row[3] in {'需重新复核', '待处理', '处理中'} for row in rows)
    pending_states = {'需重新复核', '待处理', '处理中'}
    focus = [row for row in rows if row[3] in pending_states and row[2] != '未变化']
    tasks = {'focus': len(focus), 'pending': attention,
             'firstPending': sum(row[2] == '首次核对' and row[3] in pending_states for row in rows),
             'newPending': sum(row[2] == '新增' and row[3] in pending_states for row in rows),
             'recheck': sum(row[3] == '需重新复核' for row in rows),
             'ongoing': sum(row[2] == '未变化' and row[3] in pending_states for row in rows),
             'carriedResolved': sum(row[2] == '未变化' and row[3] == '已解决' for row in rows),
             'absent': absent}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '汇总'
    overview = [['基础对账 · 先看需要处理的事项', VERSION],
                ['本次待办', attention], ['需重新复核', sum(row[3] == '需重新复核' for row in rows)],
                ['新增待办', sum(row[2] == '新增' and row[3] != '无需处理' for row in rows)],
                ['意见沿用', sum(row[2] == '未变化' and bool(row[4]) for row in rows)],
                ['下一步', '先看本次重点，点击处理入口到待处理页填写黄色列；未变化但未解决的事项仍在跟进清单。'],
                ['统计口径', '生成时统计；手动修改处理状态后，汇总数字不实时更新，下次重跑重新统计。'],
                ['左表记录数', summary['leftRows']],
                ['右表记录数', summary['rightRows']], ['左表金额合计', summary['leftTotal']],
                ['右表金额合计', summary['rightTotal']], ['使用方式', '填写黄色列；已解决记录仍保留，可用处理状态筛选。下次把填写过的工作簿交给AI。'],
                ['重跑规则', '内容未变化才沿用状态；变化时保留意见并要求重新复核。'],
                ['历史记录', '本次未出现仅代表不在本次输入中，不代表问题已解决。']]
    overview += [['金额差异怎么查', '对照左右金额和来源行回查单据，不自动断言差额原因。'],
                 ['单边记录怎么查', '核实遗漏、期间和编号口径。'],
                 ['重复 / 无效值怎么查', '先确认明细粒度与原始数据，再重新核对。']]
    overview += [['本次重点', tasks['focus']], ['旧事项仍待跟进', tasks['ongoing']],
                 ['已解决沿用', tasks['carriedResolved']], ['上次记录本次未出现', tasks['absent']],
                 ['意见填写位置', '只在待处理或已匹配页填写；本次重点为只读视图，避免两处意见不一致。']]
    overview += list(summary['counts'].items()) + list(changes.items())
    overview += source_locations(inputs)
    for number,row in enumerate(overview,1):
        for col,value in enumerate(row,1):
            text_cell(ws.cell(number,col),value)
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 85
    for row in ws:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws['A1'].font = Font(size=16, bold=True, color='17365D')
    ws.column_dimensions['A'].width = 38
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 30
    for number in range(2, 6):
        ws.cell(number, 2).font = Font(size=16, bold=True, color='17365D')
        ws.cell(number, 2).fill = PatternFill('solid', fgColor='E8F1FA')
    for number in range(6, ws.max_row + 1):
        ws.row_dimensions[number].height = 32
    review = [r for r in rows if r[1] != '一致' or r[3] == '需重新复核']
    make_focus(wb, focus, {row[10]: number for number, row in enumerate(review, 2)})
    make_sheet(wb, '待处理', review)
    make_sheet(wb, '已匹配', [r for r in rows if r[1] == '一致' and r[3] != '需重新复核'])
    make_sheet(wb, '历史记录', history)
    meta = wb.create_sheet('_meta')
    meta.append(['format', 'finance-review-v1'])
    meta.append(['profile', profile_hash])
    meta.sheet_state = 'hidden'
    wb.save(dest)
    save_followup(rows, profile, inputs, dest.with_name('跟进清单.xlsx'))
    return changes, tasks


def main():
    parser = argparse.ArgumentParser(description='保存核对配置，换表重跑，输出可填写复核意见的Excel')
    for name in ['config', 'save-config', 'left', 'right', 'out', 'previous', 'left-keys', 'right-keys', 'left-amount', 'right-amount', 'left-sheet', 'right-sheet', 'left-header-row', 'right-header-row']:
        parser.add_argument('--'+name)
    parser.add_argument('--single-key-confirmed', action='store_true')
    args = vars(parser.parse_args())
    for key in ['left', 'right', 'out']:
        if not args[key]:
            raise ValueError('缺少 --'+key)
    if bool(args['config']) == bool(args['save_config']):
        raise ValueError('首次用 --save-config；以后用 --config，二选一')
    fields = ['left_keys', 'right_keys', 'left_amount', 'right_amount', 'left_sheet', 'right_sheet', 'left_header_row', 'right_header_row', 'single_key_confirmed']
    if args['config']:
        if any(args[k] for k in fields):
            raise ValueError('复用配置时不能同时改映射；格式变化请创建新配置')
        profile = read_json(args['config'])
        if set(profile) != {'version', 'mapping', 'headers'} or profile['version'] != 1 or not set(profile['mapping']).issubset(fields):
            raise ValueError('配置格式不支持')
        mapping = profile['mapping']
    else:
        if Path(args['save_config']).exists():
            raise ValueError('配置文件已存在，请复用或另起名称')
        mapping = {k: args[k] for k in fields if args[k]}
    cli = [sys.executable, '-X', 'utf8', str(ROOT/'reconcile.py'), '--left', args['left'], '--right', args['right']]
    for key, value in mapping.items():
        if key == 'single_key_confirmed':
            if value is True: cli.append('--single-key-confirmed')
        elif isinstance(value, str):
            cli += ['--'+key.replace('_','-'), value]
        else:
            raise ValueError('配置字段类型无效')
    def execute(extra):
        result = subprocess.run(cli + extra, capture_output=True, text=True, encoding='utf-8', timeout=60)
        if result.returncode:
            raise ValueError(result.stderr.strip())
        return json.loads(result.stdout)
    input_hashes = {side: file_hash(args[side], 10 * 1024 * 1024) for side in ['left', 'right']}
    inspection = execute(['--inspect'])
    headers = {k: inspection[k]['headers'] for k in ['left', 'right']}
    if args['config'] and headers != profile['headers']:
        raise ValueError('表头与保存配置不同，请先确认字段变化并建立新配置')
    if not args['config']:
        profile = {'version': 1, 'mapping': mapping, 'headers': headers}
    out = Path(args['out'])
    if out.exists():
        raise ValueError('结果目录已存在，请使用新目录')
    # Stage outputs; invalid previous reviews never leave a partial success directory.
    with tempfile.TemporaryDirectory(dir=out.parent) as staging:
        stage = Path(staging)/'result'
        summary = execute(['--out', str(stage)])
        manifest = read_json(stage/'运行清单.json')
        if any(manifest['inputs'][side]['sha256'] != input_hashes[side] for side in input_hashes):
            raise ValueError('检查表头后输入文件发生变化，请使用稳定的导出文件重新核对')
        previous_hash = file_hash(args['previous'], 10 * 1024 * 1024) if args['previous'] else None
        changes, tasks = build_workbook(read_json(stage/'核对记录.json'), summary, profile, args['previous'], stage/'复核工作簿.xlsx', manifest['inputs'])
        manifest['workflowVersion'] = VERSION
        manifest['profileSHA256'] = digest(profile)
        manifest['reviewChanges'] = changes
        manifest['reviewTasks'] = tasks
        if args['previous']:
            if file_hash(args['previous'], 10 * 1024 * 1024) != previous_hash:
                raise ValueError('上次复核工作簿在读取后发生变化，请重新核对')
            manifest['previousReviewSHA256'] = previous_hash
        for filename in ['复核工作簿.xlsx', '跟进清单.xlsx']:
            manifest['outputs'].append({'file': filename, 'sha256': hashlib.sha256((stage/filename).read_bytes()).hexdigest()})
        message = (f"本次重点 {tasks['focus']} 项：首次待核实 {tasks['firstPending']}，新增待办 {tasks['newPending']}，需重新复核 {tasks['recheck']}。\n"
                   f"旧事项仍待跟进 {tasks['ongoing']} 项；已解决沿用 {tasks['carriedResolved']} 项。\n"
                   f"上次记录本次未出现 {tasks['absent']} 项，不代表已解决。\n"
                   "先查看复核工作簿的本次重点，点击处理入口填写原记录。跟进清单含所有未解决事项。\n"
                   "以上为生成时快照，手工修改状态后需重新运行更新。\n")
        (stage/'本次待办摘要.txt').write_text(message, encoding='utf-8')
        manifest['outputs'].append({'file': '本次待办摘要.txt', 'sha256': hashlib.sha256((stage/'本次待办摘要.txt').read_bytes()).hexdigest()})
        (stage/'运行清单.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        if any(file_hash(args[side], 10 * 1024 * 1024) != input_hashes[side] for side in input_hashes):
            raise ValueError('生成复核工作簿时输入文件发生变化，请重新核对')
        if args['save_config']:
            with Path(args['save_config']).open('x', encoding='utf-8') as f:
                json.dump(profile, f, ensure_ascii=False, indent=2)
        stage.rename(out)
    print(json.dumps({'workbook': str(out/'复核工作簿.xlsx'), 'followup': str(out/'跟进清单.xlsx'),
                      'taskSummary': str(out/'本次待办摘要.txt'), 'tasks': tasks, 'changes': changes}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print('处理失败：'+str(error), file=sys.stderr)
        sys.exit(1)
